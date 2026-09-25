"""PortableConverterMD 转换引擎测试

覆盖：各格式转换、输出目录创建、命名规则与同名去重、不支持类型报错、空 PDF 报错。
"""
import os

import pytest

from converter import convert_file
import engine


# ------------------------------------------------------------ 素材

def make_text_pdf(path: str, pages: int = 2, lines: int = 5) -> str:
    """手写一个带文本层的多页 PDF（不依赖任何 PDF 生成库）。"""
    page_nums, content_nums = [], []
    num = 4
    for _ in range(pages):
        page_nums.append(num); num += 1
        content_nums.append(num); num += 1
    objs = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: ("<< /Type /Pages /Kids [" + " ".join(f"{n} 0 R" for n in page_nums) +
            f"] /Count {pages} >>").encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for i in range(pages):
        body = "BT /F1 11 Tf 60 740 Td 14 TL\n"
        for j in range(lines):
            body += f"(Page {i+1} line {j+1}: conversion engine test text) Tj T*\n"
        body += "ET\n"
        raw = body.encode("latin-1")
        objs[content_nums[i]] = (b"<< /Length " + str(len(raw)).encode() + b" >>\nstream\n" + raw + b"endstream")
        objs[page_nums[i]] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                              f"/Resources << /Font << /F1 3 0 R >> >> "
                              f"/Contents {content_nums[i]} 0 R >>").encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in range(1, num):
        offsets[n] = len(out)
        out += f"{n} 0 obj\n".encode() + objs[n] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {num}\n".encode() + b"0000000000 65535 f \n"
    for n in range(1, num):
        out += f"{offsets[n]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {num} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(bytes(out))
    return path


def make_blank_pdf(path: str) -> str:
    """一页纯空白、无文本层的 PDF（用于验证“识别不到文字”的报错）。"""
    objs = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>",
    }
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for n in (1, 2, 3):
        offsets[n] = len(out)
        out += f"{n} 0 obj\n".encode() + objs[n] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 4\n0000000000 65535 f \n"
    for n in (1, 2, 3):
        out += f"{offsets[n]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(bytes(out))
    return path


# ------------------------------------------------------------ 基础行为

def test_convert_txt_to_markdown(tmp_path):
    src = tmp_path / "hello.txt"
    src.write_text("Hello World", encoding="utf-8")
    out = convert_file(str(src), str(tmp_path / "md_output"))
    assert os.path.exists(out) and out.endswith(".md")
    assert "Hello World" in open(out, encoding="utf-8").read()


def test_output_dir_is_created(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("content", encoding="utf-8")
    out_dir = tmp_path / "nonexistent" / "md_output"
    out = convert_file(str(src), str(out_dir))
    assert out_dir.is_dir() and os.path.exists(out)


def test_output_name_matches_input_stem(tmp_path):
    src = tmp_path / "myfile.txt"
    src.write_text("test", encoding="utf-8")
    out = convert_file(str(src), str(tmp_path / "md_output"))
    assert os.path.basename(out) == "myfile.md"


def test_same_stem_does_not_overwrite(tmp_path):
    """同名不同扩展名的输入不能互相覆盖（旧实现会静默覆盖）。"""
    d = tmp_path / "src"
    d.mkdir()
    (d / "report.txt").write_text("TXT VERSION", encoding="utf-8")
    (d / "report.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    out_dir = tmp_path / "md_output"

    first = convert_file(str(d / "report.txt"), str(out_dir))
    second = convert_file(str(d / "report.csv"), str(out_dir))

    assert first != second
    assert os.path.basename(first) == "report.md"
    assert os.path.basename(second) == "report (2).md"
    assert "TXT VERSION" in open(first, encoding="utf-8").read()
    assert "| a | b |" in open(second, encoding="utf-8").read()


def test_unsupported_extension_raises_clear_error(tmp_path):
    src = tmp_path / "weird.xyz"
    src.write_bytes(b"\x00\x01")
    with pytest.raises(RuntimeError, match="暂不支持的文件类型"):
        convert_file(str(src), str(tmp_path / "md_output"))


def test_supported_ext_has_no_audio_or_cloud(tmp_path):
    """音频/云文档智能已从支持范围移除。"""
    for ext in (".mp3", ".wav", ".doc", ".xls", ".ppt"):
        assert ext not in engine.SUPPORTED_EXT
    for ext in (".txt", ".csv", ".json", ".xml", ".html", ".docx", ".xlsx",
                ".pptx", ".pdf", ".png", ".zip"):
        assert ext in engine.SUPPORTED_EXT


# ------------------------------------------------------------ 各格式

def test_csv_becomes_table(tmp_path):
    src = tmp_path / "t.csv"
    src.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    out = convert_file(str(src), str(tmp_path / "o"))
    text = open(out, encoding="utf-8").read()
    assert "| a | b |" in text and "| 3 | 4 |" in text


def test_json_becomes_code_block(tmp_path):
    src = tmp_path / "t.json"
    src.write_text('{"k": "值"}', encoding="utf-8")
    out = convert_file(str(src), str(tmp_path / "o"))
    text = open(out, encoding="utf-8").read()
    assert text.startswith("```json") and "值" in text


def test_html_becomes_markdown(tmp_path):
    src = tmp_path / "t.html"
    src.write_text("<html><body><h1>标题</h1><p>段落</p></body></html>", encoding="utf-8")
    out = convert_file(str(src), str(tmp_path / "o"))
    text = open(out, encoding="utf-8").read()
    assert "# 标题" in text and "段落" in text


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "列A"
    table.rows[0].cells[1].text = "列B"
    src = tmp_path / "t.docx"
    doc.save(str(src))

    text = open(convert_file(str(src), str(tmp_path / "o")), encoding="utf-8").read()
    assert "第一段" in text
    assert "| 列A | 列B |" in text


def test_xlsx_sheets_and_rows(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "数量"])
    ws.append(["甲", 12])
    src = tmp_path / "t.xlsx"
    wb.save(str(src))

    text = open(convert_file(str(src), str(tmp_path / "o")), encoding="utf-8").read()
    assert "# 数据" in text
    assert "| 名称 | 数量 |" in text and "| 甲 | 12 |" in text


def test_pptx_slides(tmp_path):
    from pptx import Presentation

    pr = Presentation()
    slide = pr.slides.add_slide(pr.slide_layouts[1])
    slide.shapes.title.text = "标题页"
    slide.placeholders[1].text = "要点一"
    src = tmp_path / "t.pptx"
    pr.save(str(src))

    text = open(convert_file(str(src), str(tmp_path / "o")), encoding="utf-8").read()
    assert "## Slide 1" in text and "标题页" in text and "要点一" in text


def test_pdf_text_layer_is_extracted_without_ocr(tmp_path):
    src = make_text_pdf(str(tmp_path / "t.pdf"), pages=2)
    text = open(convert_file(src, str(tmp_path / "o")), encoding="utf-8").read()
    assert "Page 1 line 1" in text and "Page 2 line 1" in text
    assert "---" in text                      # 页分隔


def test_blank_pdf_reports_no_text(tmp_path):
    """空白页 PDF 应给出明确错误，而不是写出一个空文件。"""
    src = make_blank_pdf(str(tmp_path / "blank.pdf"))
    with pytest.raises(RuntimeError, match="未能从该 PDF 中提取到文字"):
        convert_file(src, str(tmp_path / "o"))


def test_zip_entries_are_converted(tmp_path):
    import zipfile

    src = tmp_path / "pack.zip"
    with zipfile.ZipFile(src, "w") as z:
        z.writestr("inner/note.txt", "压缩包里的内容")
        z.writestr("inner/skip.bin", b"\x00\x01")
    text = open(convert_file(str(src), str(tmp_path / "o")), encoding="utf-8").read()
    assert "## inner/note.txt" in text and "压缩包里的内容" in text
    assert "skip.bin" not in text
