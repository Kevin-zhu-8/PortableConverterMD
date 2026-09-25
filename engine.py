"""格式引擎：扩展名分发 + 每格式懒加载

相对旧实现（markitdown）的差别：
- 不做 ML 文件类型嗅探，按扩展名分发，去掉 magika/onnxruntime
- 每个格式的第三方库都在函数内部 import：用到哪个格式才付哪个格式的导入成本
- PDF 先取文本层，只对无文本页渲染 + OCR，并按块并行，内存可控
- 图片走 OCR（旧实现输出空文件）
- 输出同名不覆盖，自动加序号

引擎层不依赖 PySide6，可被 GUI / CLI / 测试直接复用。
"""
from __future__ import annotations

import csv as _csv
import json as _json
import os
import zipfile
from pathlib import Path

import ocr

# ------------------------------------------------------------ 基础格式

PDF_TEXT_MIN = 40          # 判定“该页有文本层”的最小字符数
PDF_RENDER_SCALE = 2.0     # 扫描页渲染倍率（实测 2 倍是清晰度/速度的平衡点）
ZIP_MAX_ENTRIES = 500
ZIP_MAX_ENTRY_BYTES = 50 * 1024 * 1024


def _txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _csv_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        rows = list(_csv.reader(f))
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    width = max(len(head), max((len(r) for r in body), default=0))
    head = head + [""] * (width - len(head))
    out = ["| " + " | ".join(head) + " |",
           "| " + " | ".join("---" for _ in range(width)) + " |"]
    for row in body:
        row = row + [""] * (width - len(row))
        out.append("| " + " | ".join(c.replace("|", "\\|") for c in row) + " |")
    return "\n".join(out)


def _json_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        data = _json.load(f)
    return "```json\n" + _json.dumps(data, ensure_ascii=False, indent=2) + "\n```"


def _xml_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return "```xml\n" + f.read().strip() + "\n```"


def _html_md(path: str) -> str:
    import markdownify

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return markdownify.markdownify(f.read(), heading_style="ATX").strip()


# ------------------------------------------------------------ Office

def _docx_md(path: str) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            parts.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                parts.append("| " + " | ".join(c.text.replace("\n", " ") for c in row.cells) + " |")
            parts.append("")
    return "\n\n".join(parts).strip()


def _xlsx_md(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).replace("\n", " ") for c in row]
                if any(cells):
                    parts.append("| " + " | ".join(cells) + " |")
            parts.append("")
        return "\n".join(parts).strip()
    finally:
        wb.close()


def _pptx_md(path: str) -> str:
    from pptx import Presentation

    pr = Presentation(path)
    parts: list[str] = []
    for i, slide in enumerate(pr.slides, 1):
        parts.append(f"## Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(f"> 备注：{slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n\n".join(parts).strip()


# ------------------------------------------------------------ PDF / 图片

def _pdf_md(path: str) -> str:
    """文本层优先；只对无文本页渲染 + OCR；按块并行，避免一次性渲染整本扫描件。"""
    from concurrent.futures import ThreadPoolExecutor

    import pypdfium2 as pdfium

    workers = max(1, min(4, (os.cpu_count() or 2)))
    doc = pdfium.PdfDocument(path)
    try:
        texts: dict[int, str] = {}
        todo: list[int] = []
        for i in range(len(doc)):
            tp = doc[i].get_textpage()
            try:
                page_text = (tp.get_text_range() or "").strip()
            finally:
                tp.close()
            if len(page_text) >= PDF_TEXT_MIN:
                texts[i] = page_text
            else:
                todo.append(i)

        chunk = max(1, workers * 2)
        for start in range(0, len(todo), chunk):
            batch = todo[start:start + chunk]
            bitmaps = {i: doc[i].render(scale=PDF_RENDER_SCALE) for i in batch}
            images = {i: bitmaps[i].to_pil() for i in batch}
            with ThreadPoolExecutor(max_workers=min(workers, len(batch))) as ex:
                for i, page_text in zip(batch, ex.map(ocr.recognize, images.values())):
                    texts[i] = (page_text or "").strip()
            del bitmaps, images

        if not any(texts.values()):
            raise RuntimeError("未能从该 PDF 中提取到文字（空白页、加密文档或图片无法识别）")
        return "\n\n---\n\n".join(texts[i] for i in sorted(texts) if texts[i])
    finally:
        doc.close()


def _image_md(path: str) -> str:
    from PIL import Image

    with Image.open(path) as img:
        return ocr.recognize(img)


def _zip_md(path: str) -> str:
    import tempfile

    parts: list[str] = []
    with zipfile.ZipFile(path) as z, tempfile.TemporaryDirectory() as td:
        root = os.path.realpath(td)
        entries = [i for i in z.infolist() if not i.is_dir()][:ZIP_MAX_ENTRIES]
        for info in entries:
            ext = Path(info.filename).suffix.lower()
            if ext not in _HANDLERS or info.file_size > ZIP_MAX_ENTRY_BYTES:
                continue
            target = os.path.realpath(os.path.join(td, info.filename))
            if not target.startswith(root):        # 防 zip slip
                continue
            try:
                z.extract(info, td)
                parts.append(f"## {info.filename}\n\n{_HANDLERS[ext](target)}")
            except Exception as e:                 # 单条目失败不影响整体
                parts.append(f"## {info.filename}\n\n> 转换失败：{type(e).__name__}: {e}")
    return "\n\n".join(parts).strip()


# ------------------------------------------------------------ 注册表

_HANDLERS = {
    ".txt": _txt, ".log": _txt, ".md": _txt, ".py": _txt, ".ini": _txt,
    ".yaml": _txt, ".yml": _txt, ".cfg": _txt, ".conf": _txt,
    ".csv": _csv_md, ".tsv": _csv_md,
    ".json": _json_md,
    ".xml": _xml_md, ".html": _html_md, ".htm": _html_md,
    ".docx": _docx_md, ".xlsx": _xlsx_md, ".pptx": _pptx_md,
    ".pdf": _pdf_md,
    ".png": _image_md, ".jpg": _image_md, ".jpeg": _image_md, ".bmp": _image_md,
    ".tif": _image_md, ".tiff": _image_md, ".webp": _image_md, ".gif": _image_md,
    ".zip": _zip_md,
}

_OCR_REQUIRED = {".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif"}

SUPPORTED_EXT = tuple(sorted(_HANDLERS))


def supports(path: str) -> bool:
    return Path(path).suffix.lower() in _HANDLERS


def convert_to_text(file_path: str) -> str:
    """把文件转成 Markdown 文本（不落盘）。"""
    ext = Path(file_path).suffix.lower()
    handler = _HANDLERS.get(ext)
    if handler is None:
        raise RuntimeError(f"暂不支持的文件类型：{ext or '(无扩展名)'}")

    text = handler(str(file_path)) or ""
    if not text.strip() and ext in _OCR_REQUIRED:
        raise RuntimeError("未能识别出文字（可能是空白页、加密或纯图形内容）")
    return text


def unique_path(path: str) -> str:
    """同名输出不覆盖已有文件，自动加 (2)(3)…"""
    if not os.path.exists(path):
        return path
    stem, ext = os.path.splitext(path)
    i = 2
    while os.path.exists(f"{stem} ({i}){ext}"):
        i += 1
    return f"{stem} ({i}){ext}"


def convert_file(file_path: str, output_dir: str) -> str:
    """转换单个文件并写出 .md，返回输出路径。

    先转换、再原子落盘：转换失败不会留下 0 字节文件，
    也不会把同名旧结果截断成空文件。
    """
    file_path, output_dir = str(file_path), str(output_dir)
    os.makedirs(output_dir, exist_ok=True)

    text = convert_to_text(file_path)          # 失败在此抛出，磁盘上不产生任何文件

    out_path = unique_path(os.path.join(output_dir, f"{Path(file_path).stem}.md"))
    tmp_path = out_path + ".part"
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            f.write(text)
        os.replace(tmp_path, out_path)
    except Exception:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except OSError:
            pass
        raise
    return out_path
